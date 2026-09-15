"""Excel preprocessing and transactional refresh; no HTTP/API assumptions."""
from __future__ import annotations

import json
import re
import sqlite3
from collections import Counter
from contextlib import closing
from datetime import date, datetime
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook

FIELDS = ("date", "day_of_week", "gate", "gate_name", "passage_id", "hour",
          "in_count", "out_count", "total_in", "total_out", "is_partial", "source_file")
GATES = {"자료실.정문": "front", "자료실.후문": "back"}
HOURS = range(8, 24)
WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
REQUIRED = ("수집일자", "게이트", "통로ID", "시작IN", "종료IN", "시작OUT", "종료OUT",
            "전체_IN", "전체_OUT") + tuple(f"{k}_{h:02}" for h in HOURS for k in ("IN", "OUT"))


class DataError(ValueError):
    def __init__(self, message, issues=None, code="EXCEL_FORMAT_ERROR"):
        super().__init__(message)
        self.issues = issues or []
        self.code = code

    def as_dict(self):
        # API adapter may use this exact common error envelope.
        return {"error": {"code": self.code, "message": str(self)}}


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        return date.fromisoformat(value.strip()).isoformat()
    raise ValueError("YYYY-MM-DD 또는 Excel 날짜 형식이 필요합니다.")


def _count(value):
    if isinstance(value, bool) or value is None:
        raise ValueError("빈 값/논리값은 입출입 수로 사용할 수 없습니다.")
    if isinstance(value, int) and value >= 0:
        return value
    if isinstance(value, float) and value >= 0 and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"[0-9]+", value.strip()):
        return int(value.strip())
    raise ValueError("0 이상의 정수가 필요합니다. 자동 보정하지 않습니다.")


def preprocess(path, *, sheet=None, partial_dates=None, source_name=None):
    """Return (12-field records, validation report). Any invalid row rejects file.

    partial_dates=None: conservatively mark latest date partial (unknown cutoff).
    partial_dates=[]: caller explicitly confirms every date is complete.
    """
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise DataError(".xlsx 파일만 지원합니다.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise DataError("Excel 파일을 읽을 수 없습니다.") from exc
    errors, warnings, parsed, skipped = [], [], [], []
    try:
        if sheet is None:
            if len(workbook.sheetnames) != 1:
                raise DataError("시트가 여러 개입니다. 사용할 시트 이름을 지정하세요.")
            sheet = workbook.sheetnames[0]
        if sheet not in workbook.sheetnames:
            raise DataError(f"시트를 찾을 수 없습니다: {sheet}")
        iterator = workbook[sheet].iter_rows(values_only=True)
        header = tuple(str(v).strip() if v is not None else "" for v in next(iterator, ()))
        missing = sorted(set(REQUIRED) - set(header))
        if missing or any(header.count(k) != 1 for k in REQUIRED):
            raise DataError("필수 컬럼이 없거나 중복되었습니다.", [{"missing_columns": missing}])
        for number, values in enumerate(iterator, start=2):
            if all(v is None for v in values):
                skipped.append({"row": number, "reason": "blank"})
                continue
            row = dict(zip(header, values))
            if (row.get("수집일자"), row.get("게이트"), row.get("통로ID")) in (
                ("전체", "그룹", "합계"), ("전체", "그룹", "평균")
            ):
                skipped.append({"row": number, "reason": "source_summary"})
                continue
            before = len(errors)
            try:
                day = _date(row.get("수집일자"))
            except (ValueError, TypeError) as exc:
                errors.append({"row": number, "field": "수집일자", "message": str(exc)})
                day = None
            gate_name = row.get("게이트")
            if gate_name not in GATES:
                errors.append({"row": number, "field": "게이트", "message": "알 수 없는 게이트"})
            passage = row.get("통로ID")
            if not isinstance(passage, str) or not passage.strip() or passage.startswith("="):
                errors.append({"row": number, "field": "통로ID", "message": "통로ID 문자열 필요"})
            counts = {}
            for key in REQUIRED[3:]:
                try:
                    counts[key] = _count(row.get(key))
                except ValueError as exc:
                    errors.append({"row": number, "field": key, "message": str(exc)})
            if before != len(errors):
                continue
            for kind in ("IN", "OUT"):
                hourly_sum = sum(counts[f"{kind}_{h:02}"] for h in HOURS)
                if hourly_sum != counts[f"전체_{kind}"]:
                    warnings.append({"row": number, "code": "TOTAL_HOURLY_MISMATCH", "kind": kind,
                                     "source_total": counts[f"전체_{kind}"], "hourly_sum": hourly_sum})
            parsed.append((number, day, gate_name, passage.strip(), counts))
    finally:
        workbook.close()
    duplicate_keys = Counter((day, gate) for _, day, gate, _, _ in parsed)
    for number, day, gate, _, _ in parsed:
        if duplicate_keys[(day, gate)] > 1:
            errors.append({"row": number, "field": "date+gate+hour", "message": "중복 날짜·게이트. 자동 합산하지 않습니다."})
    if errors:
        raise DataError("검증 오류로 파일 전체 반영을 중단했습니다.", errors)
    if not parsed:
        raise DataError("반영할 날짜별 데이터가 없습니다.")
    days = {p[1] for p in parsed}
    if partial_dates is None:
        partial = {max(days)}
        warnings.append({"code": "PARTIAL_DATE_UNCONFIRMED", "date": max(days),
                         "message": "수집 완료 여부 미확인: 최신 날짜를 보수적으로 부분 데이터 표시"})
    else:
        try:
            partial = {_date(d) for d in partial_dates}
        except (ValueError, TypeError) as exc:
            raise DataError("부분 데이터 날짜 형식이 잘못되었습니다.") from exc
        if not partial <= days:
            raise DataError("지정한 부분 데이터 날짜가 파일에 없습니다.")
    name = source_name or path.name
    if not isinstance(name, str) or not name.strip() or "/" in name or "\\" in name:
        raise DataError("source_name은 경로가 아닌 원본 파일명이어야 합니다.")
    records = []
    for _, day, gate_name, passage, counts in parsed:
        for hour in HOURS:
            records.append(dict(zip(FIELDS, (
                day, WEEKDAYS[date.fromisoformat(day).weekday()], GATES[gate_name], gate_name,
                passage, hour, counts[f"IN_{hour:02}"], counts[f"OUT_{hour:02}"],
                counts["전체_IN"], counts["전체_OUT"], day in partial, name))))
    records.sort(key=lambda r: (r["date"], r["gate"], r["hour"]))
    report = {"source_file": name, "sheet": sheet, "source_rows": len(parsed),
              "output_rows": len(records), "date_min": min(days), "date_max": max(days),
              "dates": len(days), "partial_dates": sorted(partial), "skipped_rows": skipped,
              "warning_counts": dict(Counter(w["code"] for w in warnings)), "warnings": warnings}
    return records, report


def refresh(path, database, *, rebuild: Callable | None = None, **options):
    """Replace incoming date+gate groups atomically; preserve all other groups.

    rebuild(all_records) returns JSON-serializable derived data, never writes to
    external systems. If it fails, previous records AND derived data stay intact.
    """
    records, report = preprocess(path, **options)
    database = Path(database)
    database.parent.mkdir(parents=True, exist_ok=True)
    try:
        with closing(sqlite3.connect(database, timeout=30)) as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                conn.execute("CREATE TABLE IF NOT EXISTS records (date TEXT, gate TEXT, hour INTEGER, payload TEXT NOT NULL, PRIMARY KEY(date,gate,hour))")
                conn.execute("CREATE TABLE IF NOT EXISTS state (id INTEGER PRIMARY KEY CHECK(id=1), payload TEXT NOT NULL)")
                for r in records:
                    old = conn.execute("SELECT payload FROM records WHERE date=? AND gate=? AND hour=?",
                                       (r["date"], r["gate"], r["hour"])).fetchone()
                    if old is not None and not json.loads(old[0])["is_partial"] and r["is_partial"]:
                        raise ValueError("완료된 날짜를 부분 데이터로 덮어쓸 수 없습니다.")
                for day, gate in sorted({(r["date"], r["gate"]) for r in records}):
                    conn.execute("DELETE FROM records WHERE date=? AND gate=?", (day, gate))
                conn.executemany("INSERT INTO records VALUES (?,?,?,?)", [
                    (r["date"], r["gate"], r["hour"], json.dumps(r, ensure_ascii=False)) for r in records])
                current = [json.loads(r[0]) for r in conn.execute("SELECT payload FROM records ORDER BY date,gate,hour")]
                derived = rebuild(current) if rebuild is not None else None
                state = {"contract_version": "v1", "record_count": len(current),
                         "integration_status": "rebuilt" if rebuild is not None else "pending_downstream",
                         "derived": derived, "last_import": report}
                conn.execute("INSERT OR REPLACE INTO state VALUES (1,?)", (json.dumps(state, ensure_ascii=False, allow_nan=False),))
                conn.commit()
            except Exception:
                conn.rollback()
                raise
    except Exception as exc:
        raise DataError("갱신 실패: 기존 데이터는 유지됩니다.", code="PROCESSING_ERROR") from exc
    return state


def read_current(database):
    """Return consistent records/state snapshot. No database creation on reads."""
    uri = Path(database).resolve().as_uri() + "?mode=ro"
    try:
        with closing(sqlite3.connect(uri, uri=True)) as conn:
            conn.execute("BEGIN")
            records = [json.loads(r[0]) for r in conn.execute("SELECT payload FROM records ORDER BY date,gate,hour")]
            row = conn.execute("SELECT payload FROM state WHERE id=1").fetchone()
            if row is None:
                raise ValueError("아직 갱신되지 않은 DB")
            return records, json.loads(row[0])
    except (sqlite3.Error, ValueError) as exc:
        raise DataError("저장된 데이터를 읽을 수 없습니다.", code="DATA_NOT_FOUND") from exc
